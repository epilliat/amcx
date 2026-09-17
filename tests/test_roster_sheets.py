"""Classeur à plusieurs onglets : c'est l'utilisateur qui désigne l'onglet.

Un `.xlsx` peut porter plusieurs promotions — le classeur d'où vient ce besoin
en a cinq, dont un groupe anglophone (39 étudiants) et un groupe francophone
(165). `openpyxl` ouvre `wb.active`, **l'onglet sélectionné au dernier
enregistrement du fichier** : AMCx chargeait donc une liste sans dire laquelle,
et les 37 copies scannées du groupe anglophone ne se rattachaient à personne.

    .venv/bin/python -m unittest discover -s tests -v
"""

import os
import sys
import tempfile
import unittest
from pathlib import Path

_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_ROOT / "auto_grading"))

import openpyxl                        # noqa: E402
from grade_imports import list_sheets, read_table   # noqa: E402
import student_list as sl              # noqa: E402


def make_book(path, sheets: dict, active: str):
    """Classeur `{onglet: [lignes]}`, avec l'onglet actif imposé."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(list(r))
    wb.active = wb.sheetnames.index(active)
    wb.save(path)


HEAD = ("nom", "prenom", "mail", "id")
EN = [HEAD, ("BADETS", "Robin", "r.b@x.fr", 3017), ("ZARPAS", "Alexandre", "a.z@x.fr", 3000),
      ("BARBEAU", "Faustine", "f.b@x.fr", 3058)]
FR = [HEAD, ("ABIDELLI", "Yazid", "y.a@x.fr", 2995), ("AJAS", "Tanguy", "t.a@x.fr", 2720)]


class SheetCase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.book = Path(self._tmp.name) / "Liste.xlsx"
        # ⚠ L'onglet actif est le FR : c'est exactement le piège — l'onglet
        # voulu (EN) n'est ni le premier, ni celui qu'openpyxl ouvre.
        make_book(self.book, {"Feuil1": [], "Reg-EN": EN, "Reg-FR": FR}, active="Reg-FR")

    def tearDown(self):
        self._tmp.cleanup()


class TestReadTable(SheetCase):
    def test_les_onglets_sont_enumerables(self):
        self.assertEqual(list_sheets(self.book), ["Feuil1", "Reg-EN", "Reg-FR"])

    def test_un_csv_n_a_pas_d_onglet(self):
        csv = Path(self._tmp.name) / "x.csv"
        csv.write_text("a,b\n1,2\n", encoding="utf-8")
        self.assertEqual(list_sheets(csv), [])

    def test_sans_onglet_demande_c_est_l_actif_pas_le_premier(self):
        """Le comportement historique, fixé ici pour qu'il ne surprenne plus."""
        rows = read_table(self.book)
        self.assertEqual(rows[1][0], "ABIDELLI")        # Reg-FR, l'onglet actif

    def test_l_onglet_demande_est_celui_qui_est_lu(self):
        rows = read_table(self.book, "Reg-EN")
        self.assertEqual([r[0] for r in rows[1:]], ["BADETS", "ZARPAS", "BARBEAU"])

    def test_un_onglet_disparu_est_une_erreur_pas_un_repli(self):
        """⚠ Retomber sur `wb.active` rendrait une tout autre promotion — le
        défaut même que le choix d'onglet corrige."""
        with self.assertRaises(ValueError) as cm:
            read_table(self.book, "Reg-DE")
        self.assertIn("Reg-DE", str(cm.exception))
        self.assertIn("Reg-EN", str(cm.exception))      # dit ce qui existe


class TestAnalyseParOnglet(SheetCase):
    def test_l_analyse_porte_sur_l_onglet_demande(self):
        a = sl.analyze_roster(self.book, "Reg-EN")
        self.assertEqual(a["sheet"], "Reg-EN")
        self.assertEqual(a["nrow"], 4)
        self.assertEqual(a["suggested"]["id_idx"], 3)
        self.assertEqual(a["suggested"]["nom_idx"], 0)

    def test_le_recap_compte_les_etudiants_de_chaque_onglet(self):
        """Des noms d'onglets bruts ne disent pas lequel porte la promo : c'est
        le compte par onglet qui rend le choix possible."""
        got = {s["name"]: s["n_students"] for s in sl.sheet_summaries(self.book)}
        self.assertEqual(got, {"Feuil1": None, "Reg-EN": 3, "Reg-FR": 2})

    def test_un_onglet_incomprehensible_ne_fait_pas_echouer_le_recap(self):
        """`None`, pas `0` : on ne sait pas ce que porte cet onglet, et annoncer
        « 0 étudiant » serait une affirmation qu'on n'a pas vérifiée. L'onglet
        reste proposé — les colonnes se choisissent à la main."""
        empty = [s for s in sl.sheet_summaries(self.book) if s["name"] == "Feuil1"][0]
        self.assertEqual((empty["nrow"], empty["n_students"]), (None, None))

    def test_au_dela_du_plafond_on_rend_les_noms_seuls(self):
        """Le récapitulatif relit le classeur une fois par onglet : ce n'est pas
        un prix à payer sans limite."""
        got = sl.sheet_summaries(self.book, max_sheets=1)
        self.assertEqual([s["name"] for s in got], ["Feuil1", "Reg-EN", "Reg-FR"])
        self.assertIsNone(got[1]["n_students"])

    def test_students_from_file_suit_l_onglet_du_mapping(self):
        cols = {"id_idx": 3, "nom_idx": 0, "prenom_idx": 1, "data_start": 1}
        en = sl.students_from_file(self.book, {**cols, "sheet": "Reg-EN"})
        fr = sl.students_from_file(self.book, {**cols, "sheet": "Reg-FR"})
        self.assertEqual([s.nom for s in en], ["BADETS", "ZARPAS", "BARBEAU"])
        self.assertEqual([s.nom for s in fr], ["ABIDELLI", "AJAS"])

    def test_un_onglet_disparu_remonte_en_RosterError(self):
        """L'UI attend une RosterError, pas une ValueError brute."""
        with self.assertRaises(sl.RosterError):
            sl.students_from_file(self.book, {"id_idx": 3, "nom_idx": 0,
                                              "data_start": 1, "sheet": "Reg-DE"})


class TestLoadStudents(SheetCase):
    """`load_students()` relit la liste à chaque démarrage : sans l'onglet en
    config, elle repartirait sur l'onglet actif — donc une autre promo."""

    def setUp(self):
        super().setUp()
        self._proj = tempfile.TemporaryDirectory()
        self._old = os.environ.get("AMCX_PROJECT_DIR")
        os.environ["AMCX_PROJECT_DIR"] = self._proj.name
        import shutil
        shutil.copy(self.book, Path(self._proj.name) / "student_list.xlsx")

    def tearDown(self):
        if self._old is None:
            os.environ.pop("AMCX_PROJECT_DIR", None)
        else:
            os.environ["AMCX_PROJECT_DIR"] = self._old
        self._proj.cleanup()
        super().tearDown()

    def _write_cfg(self, sheet):
        import json
        cfg = {"student_xlsx": "student_list.xlsx", "xlsx_id_idx": 3,
               "xlsx_nom_idx": 0, "xlsx_prenom_idx": 1, "xlsx_data_start": 1,
               "xlsx_sheet": sheet}
        (Path(self._proj.name) / "config.json").write_text(json.dumps(cfg),
                                                           encoding="utf-8")

    def test_l_onglet_en_config_est_respecte(self):
        self._write_cfg("Reg-EN")
        self.assertEqual([s.nom for s in sl.load_students()],
                         ["BADETS", "ZARPAS", "BARBEAU"])

    def test_sans_onglet_en_config_on_lit_l_actif(self):
        self._write_cfg("")
        self.assertEqual([s.nom for s in sl.load_students()], ["ABIDELLI", "AJAS"])

    def test_un_onglet_renomme_depuis_l_import_est_signale(self):
        self._write_cfg("Reg-DE")
        with self.assertRaises(sl.RosterError) as cm:
            sl.load_students()
        self.assertIn("Reg-DE", str(cm.exception))


if __name__ == "__main__":
    unittest.main()
